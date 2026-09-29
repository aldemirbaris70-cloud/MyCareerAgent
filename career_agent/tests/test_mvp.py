"""Focused end-to-end tests for the career-agent MVP."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from career_agent.analyzers.scorer import ATSScorer
from career_agent.analyzers.skills import SkillExtractor
from career_agent.core.application import CareerAgent
from career_agent.core.errors import DocumentReadError
from career_agent.core.models import MatchScore
from career_agent.generators.cover_letter import CoverLetterGenerator
from career_agent.parsers.documents import DocumentParser
from career_agent.reports.markdown import MarkdownReportWriter


class SkillAndScoreTests(unittest.TestCase):
    def test_skill_extraction_uses_canonical_names_and_boundaries(self) -> None:
        skills = SkillExtractor().extract("Python, python3, C++, and REST API")
        self.assertEqual(skills, {"Python", "C++", "REST API"})

    def test_score_reports_overlap_and_missing_skills(self) -> None:
        result = ATSScorer().score(
            "Experience Education Projects", {"Python"}, "Python SQL", {"Python", "SQL"}
        )
        self.assertEqual(result.matched_skills, ("Python",))
        self.assertEqual(result.missing_skills, ("SQL",))
        self.assertEqual(result.skill_coverage, 50.0)
        self.assertEqual(result.section_coverage, 100.0)
        self.assertEqual(result.overall, 57.5)

    def test_cover_letter_does_not_claim_unmatched_skills(self) -> None:
        score = MatchScore(0.0, 0.0, 0.0, (), ("Python",))
        letter = CoverLetterGenerator().generate("Student", "Software Intern\nPython", score)
        self.assertIn("apply for Software Intern", letter)
        self.assertNotIn("My CV highlights", letter)
        self.assertNotIn("background includes", letter)


class DocumentAndReportTests(unittest.TestCase):
    def test_docx_parser_reads_paragraphs_and_tables(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "resume.docx"
            document = Document()
            document.add_paragraph("Python developer")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "SQL project"
            document.save(path)
            text = DocumentParser().read(path)
            self.assertIn("Python developer", text)
            self.assertIn("SQL project", text)
            uploaded_text = DocumentParser().read_bytes("resume.docx", path.read_bytes())
            self.assertIn("Python developer", uploaded_text)
            self.assertIn("SQL project", uploaded_text)

    def test_uploaded_text_is_parsed_without_a_temporary_file(self) -> None:
        parser = DocumentParser()
        self.assertEqual(parser.read_bytes("cv.txt", b"Python\nSQL"), "Python\nSQL")

    def test_pdf_parser_reads_pages_and_rejects_encryption(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "resume.pdf"
            path.write_bytes(b"%PDF-1.4")
            with patch("pypdf.PdfReader") as reader_factory:
                reader_factory.return_value.is_encrypted = False
                reader_factory.return_value.pages = [unittest.mock.Mock(extract_text=lambda: "Python experience")]
                self.assertEqual(DocumentParser().read(path), "Python experience")

                reader_factory.return_value.is_encrypted = True
                with self.assertRaisesRegex(DocumentReadError, "encrypted"):
                    DocumentParser().read(path)

    def test_text_parse_pipeline_and_markdown_save(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            cv_path = root / "cv.txt"
            cv_path.write_text("Education\nProjects\nExperience\nPython\n", encoding="utf-8")
            report = CareerAgent().analyze(
                cv_path,
                job_text="Intern role requiring Python, SQL",
                candidate_name="Jordan Lee",
                company="Example Co",
            )
            output = MarkdownReportWriter().save(report, root / "reports")
            saved = output.read_text(encoding="utf-8")
            self.assertIn("ATS Compatibility", saved)
            self.assertIn("Tailored Cover Letter", saved)
            self.assertIn("SQL", saved)
            self.assertEqual(output.name, "Jordan_Lee_job_match.md")

    def test_unsupported_file_extension_is_actionable(self) -> None:
        with self.assertRaisesRegex(DocumentReadError, "Unsupported file type"):
            DocumentParser().read("resume.rtf")


if __name__ == "__main__":
    unittest.main()