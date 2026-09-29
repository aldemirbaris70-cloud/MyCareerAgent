"""Read text, PDF, and DOCX documents into plain text."""

import logging
from pathlib import Path

from career_agent.core.errors import DocumentReadError

logger = logging.getLogger(__name__)


class DocumentParser:
    """Extract text from supported document formats."""

    SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}

    def read(self, source: str | Path) -> str:
        """Read a supported document path, wrapping common failures clearly."""
        path = Path(source).expanduser()
        extension = path.suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise DocumentReadError(
                f"Unsupported file type '{extension or '(none)'}'. Use TXT, PDF, or DOCX."
            )
        if not path.is_file():
            raise DocumentReadError(f"File does not exist or is not a file: {path}")
        try:
            if extension == ".txt":
                return path.read_text(encoding="utf-8-sig")
            if extension == ".pdf":
                return self._read_pdf(path)
            return self._read_docx(path)
        except DocumentReadError:
            raise
        except Exception as error:
            logger.exception("Could not parse document %s", path)
            raise DocumentReadError(f"Could not read '{path}': {error}") from error

    @staticmethod
    def _read_pdf(path: Path) -> str:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocumentReadError(f"The PDF is encrypted: {path}")
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    @staticmethod
    def _read_docx(path: Path) -> str:
        from docx import Document

        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        table_text = [
            " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(paragraphs + [line for line in table_text if line]).strip()